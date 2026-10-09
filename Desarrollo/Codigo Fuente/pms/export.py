"""Exportaciones a Excel: plan modificado (formato original) y reporte con gráficos nativos."""
from __future__ import annotations

import datetime as _dt
import os

import numpy as np

from . import APP_NAME, VERSION
from . import config as C


def moves_rows(res):
    plan = res.plan
    names = {"semanal": "Convencional · Semanal", "diario": "Convencional · Diario", "aleatoria": "Asignación Aleatoria"}
    rm = plan.row_material()
    mats = plan.mat_names
    out = []
    for r, stage, src, dst, t in res.moves:
        d = plan.date_ord[r]
        out.append([int(plan.excel_rows[r]), int(plan.week[r]),
                    _dt.date.fromordinal(int(d)).strftime("%d/%m/%Y") if d else "", plan.phase[r], plan.poly[r],
                    int(plan.sec[r]) if float(plan.sec[r]).is_integer() else plan.sec[r],
                    mats[rm[r]] if rm[r] >= 0 else "", names.get(stage, stage), src, dst, float(t)])
    return out


MOVE_HEADERS = ["Fila Excel", "Semana", "Fecha", "Fase", "Polígono / Origen", "# Sec", "Material", "Etapa", "Desde",
                "Hacia", "Toneladas"]
REORDER_HEADERS = ["Semana", "Polígonos", "Solo Mineral", "Solo Desmonte", "Mineral y Desmonte", "Base del %",
                   "Objetivo Inician Mineral", "Inician Mineral (Plan)", "Inician Mineral (Resultado)",
                   "Inician Desmonte (Plan)", "Inician Desmonte (Resultado)", "% Inician Mineral", "Diferencia",
                   "Intercambios"]


def reorder_rows(res):
    return [[t["semana"], t["poligonos"], t["solo_mineral"], t["solo_desmonte"], t["ambos"], t["base"],
             round(t["objetivo"], 2), t["mineral_antes"], t["mineral_despues"], t["desmonte_antes"],
             t["desmonte_despues"], round(t["pct_mineral"], 2), round(t["diferencia"], 2), t["intercambios"]]
            for t in res.reorder_table]


def polygon_stack(res, use_result=True):
    """Por semana: [(polígono, tipo_inicial, t_tipo_inicial, t_otro_tipo)] en el orden del plan."""
    from .engine import row_types
    plan, cfg = res.plan, res.cfg
    types = row_types(plan, cfg)
    perm = res.perm if use_result else np.arange(plan.n)
    ym = plan.year_mask(cfg["general"].get("anio"))
    enabled = {p for p in plan.detected_phases() if cfg["fases"].get(p, True)}
    tot = plan.V[:, 0]
    out = {w: [] for w in range(1, 53)}
    for (s, e) in plan.polygon_blocks():
        w = int(plan.week[s])
        if not (1 <= w <= 52) or not ym[s] or plan.phase[s] not in enabled:
            continue
        first = None
        acc = {"M": 0.0, "D": 0.0}
        for r in range(s, e):
            t = types[perm[r]]
            if t in acc:
                first = first or t
                acc[t] += tot[perm[r]]
        if first:
            other = "D" if first == "M" else "M"
            out[w].append((plan.poly[s], first, acc[first] / 1e6, acc[other] / 1e6))
    return out


def _stack_sheet(wb, res):
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.marker import DataPoint
    from openpyxl.styles import Font
    ws = wb.create_sheet("Apilado Polígonos")
    ws["A1"] = "Material semanal por polígono (Mt)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("Primer polígono arriba. Dentro de cada polígono, el material con el que inicia va arriba "
                "(azul = Mineral, naranja = Desmonte).")
    plan_st, res_st = polygon_stack(res, False), polygon_stack(res, True)
    kmax = max([len(v) for v in plan_st.values()] + [1])
    row_tab = 60
    ws.cell(row=row_tab - 1, column=1, value="Detalle por polígono").font = Font(bold=True)
    heads = ["Semana", "Orden", "Polígono", "Inicia con (Plan)", "Inicia con (Resultado)", "Mineral (Mt)",
             "Desmonte (Mt)"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=row_tab, column=j, value=h).font = Font(bold=True)
    r = row_tab + 1
    name = {"M": "Mineral", "D": "Desmonte"}
    for w in range(1, 53):
        for k, (a, b) in enumerate(zip(plan_st[w], res_st[w]), 1):
            mt_m = a[2] if a[1] == "M" else a[3]
            mt_d = a[2] if a[1] == "D" else a[3]
            for j, v in enumerate([w, k, a[0], name[a[1]], name[b[1]], round(mt_m, 6), round(mt_d, 6)], 1):
                ws.cell(row=r, column=j, value=v)
            r += 1
    # datos de los gráficos (a la derecha del detalle)
    colors = {"M": "4472C4", "D": "ED7D31"}
    anchors = {"plan": "A4", "res": "A31"}
    base_col = 10
    for which, st in (("plan", plan_st), ("res", res_st)):
        c0 = base_col
        ws.cell(row=row_tab, column=c0, value=f"Semana ({'Plan' if which == 'plan' else 'Resultado'})")
        for w in range(1, 53):
            ws.cell(row=row_tab + w, column=c0, value=w)
        series_cols = []
        c = c0 + 1
        # orden de apilado desde abajo: último polígono primero; dentro de cada uno, el segundo tipo primero
        for k in range(kmax, 0, -1):
            for part in ("otro", "inicio"):
                ws.cell(row=row_tab, column=c, value=f"Polígono {k} · {part}")
                pts = []
                for w in range(1, 53):
                    items = st[w]
                    if k <= len(items):
                        poly, first, t1, t2 = items[k - 1]
                        val = t1 if part == "inicio" else t2
                        tipo = first if part == "inicio" else ("D" if first == "M" else "M")
                    else:
                        val, tipo = 0, "M"
                    ws.cell(row=row_tab + w, column=c, value=round(val, 6))
                    pts.append(tipo)
                series_cols.append((c, pts))
                c += 1
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "stacked"
        ch.overlap = 100
        ch.gapWidth = 60
        ch.title = f"Material por polígono: 52 semanas (Mt) · {'Plan' if which == 'plan' else 'Resultado'}"
        ch.y_axis.title = "Material (Mt)"
        ch.x_axis.title = "Semana"
        ch.legend = None
        ch.height, ch.width = 12.5, 34
        for col, pts in series_cols:
            ch.add_data(Reference(ws, min_col=col, min_row=row_tab, max_row=row_tab + 52), titles_from_data=True)
            s = ch.series[-1]
            s.graphicalProperties.solidFill = colors["M"]
            s.graphicalProperties.line.solidFill = "FFFFFF"
            s.graphicalProperties.line.width = 1200
            for i, tipo in enumerate(pts):
                if tipo != "M":
                    dp = DataPoint(idx=i)
                    dp.graphicalProperties.solidFill = colors[tipo]
                    dp.graphicalProperties.line.solidFill = "FFFFFF"
                    dp.graphicalProperties.line.width = 1200
                    s.dPt.append(dp)
        ch.set_categories(Reference(ws, min_col=c0, min_row=row_tab + 1, max_row=row_tab + 52))
        ws.add_chart(ch, anchors[which])
        base_col = c + 1


def export_plan(res, out_path, progress=None):
    """Copia del libro original con la hoja analizada actualizada (fórmulas y formato se conservan)."""
    import openpyxl
    plan = res.plan
    if not os.path.exists(plan.path):
        raise FileNotFoundError(f"No se encontró el archivo base:\n{plan.path}")
    if progress:
        progress("Abriendo libro original…", 0.05)
    wb = openpyxl.load_workbook(plan.path)
    ws = wb[plan.sheet]
    if progress:
        progress("Escribiendo plan modificado…", 0.45)
    V0, V1 = plan.V, res.V
    changed = np.abs(V1 - V0) > 1e-9
    c0 = plan.first_col + plan.content_start
    rows, cols = np.nonzero(changed)
    for r, j in zip(rows.tolist(), cols.tolist()):
        ws.cell(row=int(plan.excel_rows[r]), column=c0 + j, value=float(V1[r, j]))
    if progress:
        progress("Agregando parámetros…", 0.7)
    pws = wb.create_sheet("Parámetros " + APP_NAME[:12])
    for row in parameters_rows(res):
        pws.append(row)
    pws.column_dimensions["A"].width = 38
    pws.column_dimensions["B"].width = 60
    try:
        wb.calculation.fullCalcOnLoad = True
    except Exception:
        pass
    if progress:
        progress("Guardando…", 0.8)
    wb.save(out_path)
    return int(changed.sum())


def parameters_rows(res):
    cfg = res.cfg
    ex = cfg["experimento"]
    rows = [["Aplicación", f"{APP_NAME} v{VERSION}"], ["Escenario", cfg["general"]["nombre"]],
            ["Descripción", cfg["general"]["descripcion"]], ["Calculado el", res.when],
            ["Archivo base", res.plan.path], ["Hoja", res.plan.sheet], ["Año", cfg["general"]["anio"]],
            ["Experimento", C.EXPERIMENTS.get(ex["tipo"], ex["tipo"])]]
    if ex["tipo"] in ("reordenamiento", "integral"):
        rows += [["Polígonos que inician con mineral (%)", ex["pct_mineral"]],
                 ["Polígonos que inician con desmonte (%)", 100 - float(ex["pct_mineral"])],
                 ["Base del porcentaje", "Polígonos con mineral y desmonte" if ex["base_pct"] == "elegibles"
                  else "Todos los polígonos de la semana"]]
    if ex["tipo"] in ("aleatoria", "integral"):
        rows += [["Grado de apertura", C.APERTURAS.get(ex["apertura"], ex["apertura"])], ["Semilla", ex["semilla"]]]
    rows.append(["Fases en evaluación", ", ".join(p for p, on in cfg["fases"].items() if on)])
    for c in C.CIRCUITS:
        cc = cfg["circuitos"][c]
        info = res.circuit_info(c)
        rows += [[f"{c} · Balanceo", "Activo" if cc["activo"] else "Inactivo"],
                 [f"{c} · Receptores", ", ".join(info.receptors)], [f"{c} · Donantes", ", ".join(info.donors)],
                 [f"{c} · Materiales", ", ".join(info.materials)],
                 [f"{c} · Modo", "Objetivo Fijo" if cc["modo"] == "fijo" else "Balanceado"],
                 [f"{c} · Granularidad", cc["granularidad"].capitalize()],
                 [f"{c} · Semanal Mín / Prom / Máx (Mt)",
                  f"{cc['semanal']['min']} / {cc['semanal']['prom']} / {cc['semanal']['max']}"],
                 [f"{c} · Diario Mín / Prom / Máx (Mt)",
                  f"{cc['diario']['min']} / {cc['diario']['prom']} / {cc['diario']['max']}"]]
    return rows


def export_report(res, out_path, progress=None):
    """Reporte completo con tablas y gráficos nativos de Excel."""
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    head_fill = PatternFill("solid", fgColor="0B2545")
    head_font = Font(bold=True, color="FFFFFF")

    def sheet(title, headers, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, max(10, min(28, len(str(h)) + 4)))
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        return ws

    if progress:
        progress("Resumen…", 0.1)
    ws = wb.active
    ws.title = "Resumen"
    for row in parameters_rows(res):
        ws.append(row)
    ws.append([])
    ws.append(["Indicador", "Circuito", "Estado", "Detalle"])
    for c in ws[ws.max_row]:
        c.fill, c.font = head_fill, head_font
    for v in res.validations:
        ws.append([v["indicador"], v["circuito"], v["estado"], v["detalle"]])
        cell = ws.cell(row=ws.max_row, column=3)
        cell.font = Font(bold=True, color="1A7F37" if v["estado"] == "Ok" else "CF222E")
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 60

    if progress:
        progress("Semana por destino…", 0.3)
    dn = res.dests()
    wb_, wf = res.week_dest("base"), res.week_dest("final")
    rows = [[w + 1] + [float(wb_[w, j]) for j in range(len(dn))] + [float(wf[w, j]) for j in range(len(dn))]
            for w in range(52)]
    s = sheet("Semana por Destino", ["Semana"] + [f"Plan · {d}" for d in dn] + [f"Escenario · {d}" for d in dn], rows)
    for c in C.CIRCUITS:
        if not res.circuit_ran(c):
            continue
        info = res.circuit_info(c)
        ch = LineChart()
        ch.title = f"{c}: material semanal por receptor (t)"
        ch.y_axis.title = "t"
        ch.x_axis.title = "Semana"
        ch.height, ch.width = 9, 22
        for r in info.receptors:
            j = dn.index(r)
            ch.add_data(Reference(s, min_col=2 + len(dn) + j, min_row=1, max_row=53), titles_from_data=True)
        ch.set_categories(Reference(s, min_col=1, min_row=2, max_row=53))
        s.add_chart(ch, f"{get_column_letter(2 * len(dn) + 3)}{2 + 20 * C.CIRCUITS.index(c)}")

    if progress:
        progress("Estado semanal…", 0.45)
    st_rows = []
    for c in C.CIRCUITS:
        for r in res.status.get(c, []):
            st_rows.append([c, r["periodo"], r["receptor"], r["plan"], r["resultado"], r["min"], r["prom"], r["max"],
                            r["resultado"] - r["prom"], r["estado"]])
    sheet("Estado Semanal", ["Circuito", "Semana", "Receptor", "Plan (t)", "Escenario (t)", "Mínimo (t)",
                             "Promedio (t)", "Máximo (t)", "Diferencia vs Promedio (t)", "Estado"], st_rows)
    d_rows = []
    for c in C.CIRCUITS:
        for r in res.daily_status.get(c, []):
            d_rows.append([c, r["periodo"], r["receptor"], r["plan"], r["resultado"], r["min"], r["prom"], r["max"],
                           r["estado"]])
    if d_rows:
        sheet("Estado Diario", ["Circuito", "Fecha", "Receptor", "Plan (t)", "Escenario (t)", "Mínimo (t)",
                                "Promedio (t)", "Máximo (t)", "Estado"], d_rows)

    if progress:
        progress("Materiales…", 0.55)
    plan = res.plan
    rm = plan.row_material()
    mats = plan.mat_names
    m_rows = []
    for d in dn:
        j = plan.vcol(plan.dest_idx[plan.dest_names.index(d)])
        for k, m in enumerate(mats):
            sel = rm == k
            a, b = float(plan.V[sel, j].sum()), float(res.V_bal[sel, j].sum())
            if a > 0 or b > 0:
                m_rows.append([d, res.cfg["destinos"].get(d, {}).get("tipo", ""), m, a, b, b - a])
    sheet("Destino y Material", ["Destino", "Tipo de Destino", "Material", "Plan (t)", "Escenario (t)", "Diferencia (t)"],
          m_rows)

    if progress:
        progress("Movimientos…", 0.65)
    sheet("Movimientos", MOVE_HEADERS, moves_rows(res))

    if res.reorder_table:
        if progress:
            progress("Reordenamiento…", 0.8)
        s = sheet("Reordenamiento", REORDER_HEADERS, reorder_rows(res))
        ch = BarChart()
        ch.type = "col"
        ch.grouping = "clustered"
        ch.title = "Polígonos que inician con mineral: plan vs resultado"
        ch.height, ch.width = 9, 24
        ch.add_data(Reference(s, min_col=8, max_col=9, min_row=1, max_row=53), titles_from_data=True)
        ch.set_categories(Reference(s, min_col=1, min_row=2, max_row=53))
        s.add_chart(ch, "P2")
    if progress:
        progress("Apilado por polígono…", 0.86)
    _stack_sheet(wb, res)
    if progress:
        progress("Guardando…", 0.94)
    wb.save(out_path)
