"""Lectura del Excel de entrada y detección de la tabla principal.

La tabla empieza en la celda «# Sec». La primera fila es el encabezado y se
leen todas las filas cuyo «# Sec» no está vacío. Las columnas se identifican
por su nombre (con posiciones de respaldo) y la frontera Destinos/Materiales
se detecta porque, en cada fila, la suma de destinos y la suma de materiales
igualan al «Total Material».
"""
from __future__ import annotations

import datetime as _dt
import os
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field

import numpy as np

ROLE_LABELS = {
    "sec": "# Sec",
    "semana": "Semana",
    "fase": "Fase",
    "malla": "Malla / Stock",
    "poligono": "Polígono / Origen",
    "predecesor": "Polígono Predecesor",
    "fecha": "Fecha Liberación",
    "proceso": "Proceso de Liberación",
    "total": "Total Material",
}

# Posiciones de respaldo (1 = columna «# Sec»), según el formato de referencia.
_FALLBACK_POS = {"semana": 2, "fase": 3, "malla": 4, "poligono": 5, "predecesor": 6,
                 "fecha": 7, "proceso": 12, "total": 13}
_FALLBACK_DEST = (14, 59)
_FALLBACK_MAT = (60, 74)


class PlanError(Exception):
    pass


def norm(text) -> str:
    s = "" if text is None else str(text)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", s).strip().lower()


def is_phase_value(value) -> bool:
    """Una fase válida no es vacía ni 0 (las filas con Fase = 0 no se detectan)."""
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return value != 0
    s = str(value).strip()
    if not s:
        return False
    try:
        return float(s.replace(",", ".")) != 0
    except ValueError:
        return True


def phase_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def natural_key(s: str):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(s))]


def to_float(v) -> float:
    if v is None or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        f = float(v)
        return f if np.isfinite(f) else 0.0
    s = str(v).strip().replace(",", ".")
    if not s:
        return 0.0
    try:
        f = float(s)
        return f if np.isfinite(f) else 0.0
    except ValueError:
        return 0.0


_EPOCH = _dt.date(1899, 12, 30)


def to_date(v):
    if v is None:
        return None
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    if isinstance(v, (int, float)):
        if 1 <= v < 2958466:
            return _EPOCH + _dt.timedelta(days=int(v))
        return None
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d"):
        try:
            return _dt.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def list_sheets(path: str) -> list[str]:
    """Lista de hojas leyendo solo workbook.xml (instantáneo)."""
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("xl/workbook.xml").decode("utf-8", "replace")
        names = re.findall(r'<(?:\w+:)?sheet\b[^>]*\bname="([^"]*)"', xml)
        import html
        return [html.unescape(n) for n in names]
    except Exception:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()


@dataclass
class Plan:
    path: str
    sheet: str
    header_row: int            # fila Excel del encabezado (1-based)
    first_col: int             # columna Excel de «# Sec» (1-based)
    headers: list
    roles: dict                # rol -> índice relativo (0-based) dentro de headers
    dest_idx: list             # índices relativos de columnas destino
    mat_idx: list              # índices relativos de columnas material
    content_start: int         # índice relativo de «Total Material»
    content_end: int           # índice relativo de la última columna de material (inclusive)
    excel_rows: np.ndarray
    V: np.ndarray              # valores numéricos de content_start..content_end
    sec: np.ndarray
    week: np.ndarray
    date_ord: np.ndarray       # ordinal de fecha (0 = sin fecha)
    phase: list
    poly: list
    malla: list
    proceso: list
    detection_note: str = ""
    file_size: int = 0
    file_mtime: float = 0.0
    _cache: dict = field(default_factory=dict, repr=False)

    # ---- columnas -------------------------------------------------------------
    @property
    def n(self) -> int:
        return len(self.week)

    def vcol(self, rel_idx: int) -> int:
        """Posición dentro de V de una columna relativa."""
        return rel_idx - self.content_start

    @property
    def dest_names(self) -> list:
        return [self.headers[i] for i in self.dest_idx]

    @property
    def mat_names(self) -> list:
        return [self.headers[i] for i in self.mat_idx]

    @property
    def prop_idx(self) -> list:
        used = set(self.roles.values()) | set(self.dest_idx) | set(self.mat_idx)
        return [i for i in range(len(self.headers)) if i not in used]

    def dest_matrix(self, V=None) -> np.ndarray:
        V = self.V if V is None else V
        return V[:, [self.vcol(i) for i in self.dest_idx]]

    def mat_matrix(self, V=None) -> np.ndarray:
        V = self.V if V is None else V
        return V[:, [self.vcol(i) for i in self.mat_idx]]

    def dest_sums(self) -> dict:
        s = self.dest_matrix().sum(axis=0)
        return {self.headers[i]: float(v) for i, v in zip(self.dest_idx, s)}

    def mat_sums(self) -> dict:
        s = self.mat_matrix().sum(axis=0)
        return {self.headers[i]: float(v) for i, v in zip(self.mat_idx, s)}

    def row_material(self) -> np.ndarray:
        """Índice (en mat_idx) del material de cada fila; -1 si no tiene material."""
        if "row_mat" not in self._cache:
            M = self.mat_matrix()
            idx = np.argmax(M, axis=1) if M.shape[1] else np.zeros(self.n, int)
            ok = M.max(axis=1) > 1e-9 if M.shape[1] else np.zeros(self.n, bool)
            self._cache["row_mat"] = np.where(ok, idx, -1)
        return self._cache["row_mat"]

    def detected_phases(self) -> list:
        seen = []
        s = set()
        for p in self.phase:
            if p and p not in s:
                s.add(p)
                seen.append(p)
        return sorted(seen, key=natural_key)

    def years(self) -> dict:
        out = {}
        valid = (self.week >= 1) & (self.week <= 52) & (self.date_ord > 0)
        for o in self.date_ord[valid]:
            y = _dt.date.fromordinal(int(o)).year
            out[y] = out.get(y, 0) + 1
        return out

    def detect_year(self):
        ys = self.years()
        if not ys:
            return None
        return max(ys.items(), key=lambda kv: (kv[1], kv[0]))[0]

    def year_mask(self, year) -> np.ndarray:
        m = (self.week >= 1) & (self.week <= 52) & (self.date_ord > 0)
        if year:
            lo = _dt.date(int(year), 1, 1).toordinal()
            hi = _dt.date(int(year), 12, 31).toordinal()
            m &= (self.date_ord >= lo) & (self.date_ord <= hi)
        return m

    def polygon_blocks(self) -> list:
        """Bloques contiguos de filas con el mismo Polígono / Origen y la misma semana."""
        if "blocks" not in self._cache:
            blocks = []
            start = 0
            for i in range(1, self.n + 1):
                if i == self.n or self.poly[i] != self.poly[start] or self.week[i] != self.week[start]:
                    blocks.append((start, i))
                    start = i
            self._cache["blocks"] = blocks
        return self._cache["blocks"]


# ---------------------------------------------------------------------------
def _find_role(headers, role):
    keys = [norm(h) for h in headers]

    def first(pred):
        for i, k in enumerate(keys):
            if pred(k):
                return i
        return None

    if role == "sec":
        return first(lambda k: k.replace(" ", "") in ("#sec", "nsec", "sec"))
    if role == "semana":
        return first(lambda k: k == "semana" or k.startswith("semana") or k == "week")
    if role == "fase":
        return first(lambda k: (k in ("fase", "phase") or k.startswith("fase") or k.startswith("phase"))
                     and "predecesor" not in k)
    if role == "malla":
        return first(lambda k: "malla" in k)
    if role == "poligono":
        i = first(lambda k: "poligono" in k and "origen" in k)
        if i is None:
            i = first(lambda k: k.startswith("poligono") and "predecesor" not in k)
        return i
    if role == "predecesor":
        return first(lambda k: "predecesor" in k)
    if role == "fecha":
        return first(lambda k: "fecha" in k or k.startswith("date"))
    if role == "proceso":
        return first(lambda k: "proceso" in k or "liberacion" in k and "fecha" not in k)
    if role == "total":
        i = first(lambda k: k.startswith("total material"))
        if i is None:
            i = first(lambda k: k.startswith("total"))
        return i
    return None


def _split_dest_mat(X: np.ndarray, total: np.ndarray):
    """Devuelve (dest_end, mat_start, mat_end) relativos a X (inclusive) o None.

    X contiene las columnas posteriores a «Total Material». Los destinos son el
    menor prefijo cuya suma iguala el total en todas las filas; los materiales,
    el siguiente bloque que también lo iguala.
    """
    rows = total > 1e-6
    if rows.sum() == 0 or X.shape[1] < 2:
        return None
    Xr = X[rows]
    tr = total[rows]
    tol = np.maximum(1e-6 * np.abs(tr), 1e-4)
    cum = np.cumsum(Xr, axis=1)
    ok = (np.abs(cum - tr[:, None]) <= tol[:, None]).mean(axis=0)
    cand = np.where(ok >= 0.995)[0]
    if len(cand) == 0:
        return None
    dest_end = int(cand[0])
    colsum = np.abs(X).sum(axis=0)
    after = [j for j in range(dest_end + 1, X.shape[1]) if colsum[j] > 1e-9]
    if not after:
        return None
    mat_start = after[0]
    cum2 = np.cumsum(Xr[:, mat_start:], axis=1)
    ok2 = (np.abs(cum2 - tr[:, None]) <= tol[:, None]).mean(axis=0)
    cand2 = np.where(ok2 >= 0.995)[0]
    if len(cand2) == 0:
        return None
    mat_end = mat_start + int(cand2[0])
    # columnas vacías (todo cero) entre destinos y materiales: pertenecen a destinos
    dest_end = mat_start - 1
    return dest_end, mat_start, mat_end


def load_plan(path: str, sheet: str, progress=None) -> Plan:
    import openpyxl

    if progress:
        progress("Abriendo libro…", 0.02)
    st = os.stat(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise PlanError(f"La hoja «{sheet}» no existe en el libro.")
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        header_row = first_col = None
        headers = None
        r = 0
        for row in it:
            r += 1
            for c, v in enumerate(row):
                if v is not None and norm(v).replace(" ", "") == "#sec":
                    header_row, first_col = r, c + 1
                    break
            if header_row:
                raw = list(row[first_col - 1:])
                headers = []
                for v in raw:
                    if v is None or str(v).strip() == "":
                        break
                    headers.append(str(v).strip())
                break
            if r > 300:
                break
        if not header_row:
            raise PlanError("No se encontró la celda «# Sec» en la hoja seleccionada.")
        ncol = len(headers)
        c0 = first_col - 1
        data = []
        excel_rows = []
        est = max(ws.max_row or 0, 1)
        for row in it:
            r += 1
            if len(row) <= c0:
                continue
            vals = row[c0:c0 + ncol]
            if vals[0] is None or (isinstance(vals[0], str) and not vals[0].strip()):
                continue
            if len(vals) < ncol:
                vals = tuple(vals) + (None,) * (ncol - len(vals))
            data.append(vals)
            excel_rows.append(r)
            if progress and len(data) % 2000 == 0:
                progress(f"Leyendo filas… {len(data):,}", 0.05 + 0.75 * min(1.0, r / est))
    finally:
        wb.close()
    if not data:
        raise PlanError("La tabla no contiene filas con datos debajo de «# Sec».")
    if progress:
        progress("Detectando columnas…", 0.85)

    roles = {}
    notes = []
    for role in ROLE_LABELS:
        i = _find_role(headers, role)
        if i is None and role in _FALLBACK_POS and _FALLBACK_POS[role] <= ncol:
            i = _FALLBACK_POS[role] - 1
            notes.append(f"{ROLE_LABELS[role]} por posición")
        if i is not None:
            roles[role] = i
    roles["sec"] = 0
    for need in ("semana", "poligono", "total"):
        if need not in roles:
            raise PlanError(f"No se encontró la columna «{ROLE_LABELS[need]}».")

    t = roles["total"]
    cols = list(zip(*data))
    numeric_after = np.array([[to_float(v) for v in cols[j]] for j in range(t + 1, ncol)]).T \
        if ncol > t + 1 else np.zeros((len(data), 0))
    total = np.array([to_float(v) for v in cols[t]])
    split = _split_dest_mat(numeric_after, total)
    if split:
        dest_end, mat_start, mat_end = split
        dest_idx = list(range(t + 1, t + 1 + dest_end + 1))
        mat_idx = list(range(t + 1 + mat_start, t + 1 + mat_end + 1))
        # incluye columnas de material sin tonelaje que siguen al bloque si su
        # nombre es corto (p. ej. R1, R2) y su suma es cero
        j = mat_idx[-1] + 1
        while j < ncol and np.abs(numeric_after[:, j - t - 1]).sum() == 0 and len(headers[j]) <= 6 \
                and not any(ch in headers[j] for ch in "(%"):
            mat_idx.append(j)
            j += 1
    else:
        a, b = _FALLBACK_DEST
        c, d = _FALLBACK_MAT
        if d > ncol:
            raise PlanError("No fue posible separar Destinos y Materiales.")
        dest_idx = list(range(a - 1, b))
        mat_idx = list(range(c - 1, d))
        notes.append("Destinos/Materiales por posición")
    content_start, content_end = t, mat_idx[-1]
    V = np.zeros((len(data), content_end - content_start + 1))
    V[:, 0] = total
    for j in range(content_start + 1, content_end + 1):
        V[:, j - content_start] = numeric_after[:, j - t - 1]

    def col(role, conv):
        if role not in roles:
            return [conv(None)] * len(data)
        return [conv(v) for v in cols[roles[role]]]

    week = np.array(col("semana", lambda v: int(to_float(v))), dtype=int)
    dates = col("fecha", to_date)
    date_ord = np.array([d.toordinal() if d else 0 for d in dates], dtype=np.int64)
    phase = col("fase", lambda v: phase_text(v) if is_phase_value(v) else "")
    poly = col("poligono", lambda v: "" if v is None else str(v).strip())
    malla = col("malla", lambda v: "" if v is None else str(v).strip())
    proceso = col("proceso", lambda v: "" if v is None else str(v).strip())
    sec = np.array(col("sec", lambda v: to_float(v)), dtype=float)
    if progress:
        progress("Tabla lista", 1.0)
    return Plan(path=path, sheet=sheet, header_row=header_row, first_col=first_col, headers=headers,
                roles=roles, dest_idx=dest_idx, mat_idx=mat_idx, content_start=content_start,
                content_end=content_end, excel_rows=np.array(excel_rows, dtype=np.int64), V=V, sec=sec,
                week=week, date_ord=date_ord, phase=phase, poly=poly, malla=malla, proceso=proceso,
                detection_note="; ".join(notes), file_size=st.st_size, file_mtime=st.st_mtime)
