"""Pruebas del motor (pytest).

Usan un Excel sintético generado al vuelo (independiente de cualquier plan real)
y, si está disponible, el Excel de la carpeta de trabajo.
"""
import datetime as dt
import os

import numpy as np
import pytest

import casos  # noqa: F401  (agrega el código fuente al path)
from pms import config as C
from pms import engine
from pms.plan import list_sheets, load_plan

DESTS = ["Rec_A", "Rec_B", "Don_1", "Don_2", "Don_3", "Stock_X", "Vacio_Z"]
MATS = ["M1", "M2", "W1", "W2", "W3", "R1"]


def make_xlsx(path, seed=0, weeks=52):
    import openpyxl
    rng = np.random.default_rng(seed)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws["B1"] = "Hoja de prueba"
    head = ["# Sec", "Semana", "Fase", "Malla / Stock", "Poligono / Origen", "Polígono Predecesor", "Fecha Liberación",
            "Limite A", "Limite B", "Limite C", "Limite D", "Proceso a la Liberación", "Total Material (t)"] + \
        DESTS + MATS + ["Cresta (m)", "Ore", "Waste"]
    for j, h in enumerate(head):
        ws.cell(row=4, column=2 + j, value=h)
    r = 5
    sec = 1
    # filas de stock (semana 0, fase 0) que no deben detectarse como fase
    for k in range(5):
        vals = [sec, 0, 0, f"STK_{k}", f"STK_{k}", "N/A", None, 1, 1, 1, 1, "N/A", 0.0] + [0.0] * len(DESTS) + \
            [0.0] * len(MATS) + [0, 0, 0]
        for j, v in enumerate(vals):
            ws.cell(row=r, column=2 + j, value=v)
        r += 1
        sec += 1
    year = 2031
    for w in range(1, weeks + 1):
        npoly = rng.integers(4, 9)
        for p in range(npoly):
            phase = f"FASE {rng.integers(1, 4)}"
            poly = f"P{w:02d}_{p:02d}"
            nrows = rng.integers(1, 6)
            day = dt.datetime(year, 1, 1) + dt.timedelta(days=int((w - 1) * 7 + rng.integers(0, 7)))
            for _ in range(nrows):
                mi = int(rng.integers(0, 5))
                tons = float(rng.uniform(2_000, 60_000))
                mat = [0.0] * len(MATS)
                mat[mi] = tons
                dest = [0.0] * len(DESTS)
                pool = [0, 1, 2, 3, 4, 5]
                k = int(rng.integers(1, 4))
                chosen = rng.choice(pool, size=k, replace=False)
                wts = rng.dirichlet(np.ones(k))
                for c, wt in zip(chosen, wts):
                    dest[int(c)] = tons * float(wt)
                dest[-2] += tons - sum(dest)   # cierre exacto en Stock_X
                vals = [sec, w, phase, poly, poly, "N/A", day, 1, 4, 3, 4, "CARGUIO", tons] + dest + mat + \
                    [0, sum(mat[:2]), sum(mat[2:])]
                for j, v in enumerate(vals):
                    ws.cell(row=r, column=2 + j, value=v)
                r += 1
                sec += 1
    wb.save(path)
    return path


def synth_config(plan, experimento="convencional", modo="fijo", gran="semanal", **ex):
    cfg = C.new_config()
    cfg["general"].update({"archivo": plan.path, "hoja": plan.sheet, "nombre": "Sintético"})
    C.ensure_plan_defaults(plan, cfg)
    for m in cfg["materiales"]:
        cfg["materiales"][m] = "Mineral" if m.startswith("M") else "Desmonte"
    for d in cfg["destinos"]:
        cfg["destinos"][d] = {"tipo": "N/A", "rol": "Donante"}
    cfg["destinos"].update({"Rec_A": {"tipo": "Desmonte", "rol": "Receptor"},
                            "Rec_B": {"tipo": "Desmonte", "rol": "Receptor"},
                            "Don_1": {"tipo": "Desmonte", "rol": "Donante"},
                            "Don_2": {"tipo": "Desmonte", "rol": "Donante"},
                            "Don_3": {"tipo": "Mineral", "rol": "Donante"},
                            "Stock_X": {"tipo": "Mineral", "rol": "Receptor"}})
    C.refresh_circuits(plan, cfg)
    for c in C.CIRCUITS:
        cc = cfg["circuitos"][c]
        cc["modo"], cc["granularidad"] = modo, gran
    cfg["experimento"]["tipo"] = experimento
    cfg["experimento"].update(ex)
    return cfg


@pytest.fixture(scope="module")
def plan(tmp_path_factory):
    p = tmp_path_factory.mktemp("datos") / "sintetico.xlsx"
    make_xlsx(str(p))
    return load_plan(str(p), "Datos")


def test_deteccion(plan):
    assert list_sheets(plan.path) == ["Datos"]
    assert (plan.header_row, plan.first_col) == (4, 2)
    assert plan.dest_names == DESTS
    assert plan.mat_names[:5] == MATS[:5]
    assert set(plan.detected_phases()) == {"FASE 1", "FASE 2", "FASE 3"}     # la fase 0 no se detecta
    assert plan.detect_year() == 2031
    assert "Vacio_Z" not in C.detected_destinations(plan)                   # suma 0: no se considera
    assert "R1" not in C.detected_materials(plan)
    blocks = plan.polygon_blocks()
    assert all(len({plan.poly[i] for i in range(a, b)}) == 1 for a, b in blocks)


def _check(res, plan):
    bad = [v for v in res.validations if v["estado"] != "Ok" and not v["indicador"].startswith("Brechas")]
    assert not bad, bad
    D0, D1 = plan.dest_matrix(), plan.dest_matrix(res.V)
    np.testing.assert_allclose(D1.sum(axis=0), D0.sum(axis=0), atol=1e-4)            # columnas (anual)
    np.testing.assert_allclose(plan.mat_matrix(res.V).sum(axis=0), plan.mat_matrix().sum(axis=0), atol=1e-6)
    np.testing.assert_allclose(plan.dest_matrix(res.V_bal).sum(axis=1), D0.sum(axis=1), atol=1e-4)   # filas
    assert (D1 > -1e-7).all()


@pytest.mark.parametrize("exp", ["convencional", "reordenamiento", "aleatoria", "integral"])
@pytest.mark.parametrize("modo", ["fijo", "balanceado"])
def test_experimentos(plan, exp, modo):
    cfg = synth_config(plan, exp, modo)
    res = engine.run(plan, cfg)
    _check(res, plan)
    assert res.lp_info[("Desmonte", "semanal")]["ok"]
    assert len(res.moves) > 0


def test_mejora_brechas(plan):
    cfg = synth_config(plan)
    res = engine.run(plan, cfg)
    info = res.circuit_info("Desmonte")
    lo, de, hi = C.objective_values(cfg, "Desmonte", "semanal")
    dn = res.dests()
    before = sum(1 for w in range(52) for r in info.receptors
                 if not lo <= res.week_dest("base")[w, dn.index(r)] <= hi)
    assert res.gaps("Desmonte") < before


def test_diario(plan):
    cfg = synth_config(plan, gran="diario")
    res = engine.run(plan, cfg)
    _check(res, plan)
    assert "Desmonte" in res.daily_status
    # el refinamiento diario conserva los totales semanales del cálculo semanal
    cfg_w = synth_config(plan)
    res_w = engine.run(plan, cfg_w)
    np.testing.assert_allclose(res.week_dest("final"), res_w.week_dest("final"), atol=1e-3)


def test_matriz_bloqueada(plan):
    cfg = synth_config(plan)
    cfg["circuitos"]["Desmonte"]["matriz"] = {"Rec_A": {"W1": False}, "Don_1": {"W2": False}}
    res = engine.run(plan, cfg)
    _check(res, plan)
    rm = plan.row_material()
    for d, m in (("Rec_A", "W1"), ("Don_1", "W2")):
        sel = rm == plan.mat_names.index(m)
        j = plan.vcol(plan.dest_idx[plan.dest_names.index(d)])
        np.testing.assert_allclose(res.V_bal[sel, j], plan.V[sel, j])


def test_prioridad_exclusion(plan):
    cfg = synth_config(plan)
    for r in ("Rec_A", "Rec_B"):
        items = cfg["circuitos"]["Desmonte"]["prioridad"][r]["donantes"]["items"]
        for it in items:
            if it[0] == "Don_2":
                it[1] = False
    res = engine.run(plan, cfg)
    assert not any(mv[2] == "Don_2" and mv[3] in ("Rec_A", "Rec_B") for mv in res.moves if mv[1] == "semanal")


def test_reordenamiento(plan):
    cfg = synth_config(plan, "reordenamiento", pct_mineral=50, base_pct="elegibles")
    res = engine.run(plan, cfg)
    perm = res.perm
    # el contenido solo se mueve dentro del bloque de su polígono
    for a, b in plan.polygon_blocks():
        assert set(perm[a:b].tolist()) == set(range(a, b))
    # columnas fijas intactas (semana, fase, polígono, # Sec) y contenido permutado como unidad
    np.testing.assert_allclose(res.V, res.V_bal[perm])
    # cada polígono con ambos tipos: se cumple el reparto factible más cercano por semana
    for t in res.reorder_table:
        if t["base"]:
            assert abs(t["diferencia"]) <= 0.5 + 1e-9


def test_reordenamiento_todos(plan):
    for pct in (0, 30, 100):
        cfg = synth_config(plan, "reordenamiento", pct_mineral=pct, base_pct="todos")
        res = engine.run(plan, cfg)
        _check(res, plan)
        for t in res.reorder_table:
            assert t["solo_mineral"] <= t["mineral_despues"] <= t["solo_mineral"] + t["ambos"]


def test_aleatoria_reproducible_y_grados(plan):
    out = {}
    for ap in ("estricto", "moderado", "flexible"):
        cfg = synth_config(plan, "aleatoria", apertura=ap, semilla=7)
        a = engine.run(plan, cfg)
        b = engine.run(plan, cfg)
        np.testing.assert_allclose(a.V, b.V)
        _check(a, plan)
        info = a.circuit_info("Desmonte")
        cols = [plan.vcol(plan.dest_idx[plan.dest_names.index(d)]) for d in info.dests]
        X = a.V_bal[:, cols]
        rows = X.sum(axis=1) > 1
        out[ap] = (X[rows] > 1e-6).sum(axis=1).mean()
        # la aleatoria no altera los totales por semana y destino del Convencional
        conv = engine.run(plan, synth_config(plan))
        np.testing.assert_allclose(a.week_dest("final"), conv.week_dest("final"), atol=1e-3)
    assert out["estricto"] <= out["moderado"] <= out["flexible"]


def test_pmsx(plan, tmp_path):
    cfg = synth_config(plan, "integral")
    res = engine.run(plan, cfg)
    p = str(tmp_path / "x.pmsx")
    C.save_project(p, cfg, res)
    cfg2, stored = C.load_project(p)
    res2 = engine.Result.from_storage(*stored)
    np.testing.assert_allclose(res2.V, res.V)
    assert cfg2["experimento"]["tipo"] == "integral"
    assert res2.reorder_table == res.reorder_table


def test_exportaciones(plan, tmp_path):
    from pms.export import export_plan, export_report
    import openpyxl
    cfg = synth_config(plan, "integral")
    res = engine.run(plan, cfg)
    out = str(tmp_path / "plan.xlsx")
    n = export_plan(res, out)
    assert n > 0
    wb = openpyxl.load_workbook(out, data_only=False)
    ws = wb["Datos"]
    r = int(plan.excel_rows[10])
    c0 = plan.first_col + plan.content_start
    for j in range(res.V.shape[1]):
        assert abs(float(ws.cell(row=r, column=c0 + j).value or 0) - res.V[10, j]) < 1e-6
    export_report(res, str(tmp_path / "reporte.xlsx"))


@pytest.mark.skipif(not os.path.exists(casos.INPUT), reason="Excel de trabajo no disponible")
def test_excel_de_trabajo():
    plan = load_plan(casos.INPUT, list_sheets(casos.INPUT)[0])
    assert plan.dest_idx[0] == 13 and plan.dest_idx[-1] == 58      # columnas 14–59
    assert plan.mat_idx[0] == 59 and plan.mat_idx[-1] == 73        # columnas 60–74
    for exp in ("convencional", "integral"):
        res = engine.run(plan, casos.base_config(plan, exp))
        _check(res, plan)


def _poly_orders(res, plan):
    """Por polígono: lista de (tipo, rango) de los registros clasificados en el orden final."""
    from pms.engine import row_types
    types = row_types(plan, res.cfg)
    out = []
    for a, b in plan.polygon_blocks():
        rows = [r for r in range(a, b) if types[res.perm[r]] in ("M", "D")]
        if rows:
            out.append([res.perm[r] for r in rows])
    return out, types


@pytest.mark.parametrize("exp", ["reordenamiento", "integral"])
def test_priorizacion_destinos(plan, exp):
    cfg = synth_config(plan, exp, pct_mineral=40, base_pct="elegibles")
    pd = C.sync_dest_priority(plan, cfg)
    pd["activo"] = True
    pd["desmonte"]["items"] = [["Don_2", True], ["Rec_B", True], ["Rec_A", True], ["Don_1", True], ["Stock_X", True],
                               ["Don_3", True]]
    pd["mineral"]["items"] = [["Stock_X", True], ["Don_3", True], ["Rec_A", True], ["Rec_B", True], ["Don_1", True],
                              ["Don_2", True]]
    res = engine.run(plan, cfg)
    _check(res, plan)
    orders, types = _poly_orders(res, plan)
    V = res.V_bal
    for key, start in (("desmonte", "D"), ("mineral", "M")):
        names = [n for n, on in pd[key]["items"] if on]
        cols = [plan.vcol(plan.dest_idx[plan.dest_names.index(d)]) for d in names]

        def rank(r):
            hit = np.nonzero(V[r, cols] > 1e-9)[0]
            return int(hit[0]) if len(hit) else len(cols)
        checked = 0
        for order in orders:
            if types[order[0]] != start:
                continue
            rest = order[1:]
            # tras el primer registro (forzado al tipo de inicio), el orden respeta la prioridad
            assert [rank(r) for r in rest] == sorted(rank(r) for r in rest)
            first_needed = [r for r in order if types[r] == start]
            assert rank(order[0]) == min(rank(r) for r in first_needed)
            checked += 1
        assert checked > 0
    # sigue cumpliendo el reparto de inicios pedido
    for t in res.reorder_table:
        if t["base"]:
            assert abs(t["diferencia"]) <= 0.5 + 1e-9
    assert sum(t["filas_movidas"] for t in res.reorder_table) > 0


def test_priorizacion_inactiva_igual_a_v21(plan):
    cfg = synth_config(plan, "reordenamiento")
    a = engine.run(plan, cfg)
    C.sync_dest_priority(plan, cfg)          # listas presentes pero desactivadas
    b = engine.run(plan, cfg)
    np.testing.assert_array_equal(a.perm, b.perm)
